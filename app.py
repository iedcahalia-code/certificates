import io
import re
import zipfile
from datetime import datetime

import pandas as pd
from flask import Flask, flash, redirect, render_template, request, send_file, url_for
from pptx import Presentation

app = Flask(__name__)
app.secret_key = "bulk-certificate-secret"

PLACEHOLDER_PATTERN = re.compile(r"{{\s*([a-zA-Z0-9_]+)\s*}}")


def extract_placeholders_from_pptx(file_bytes: bytes) -> set[str]:
    placeholders: set[str] = set()
    presentation = Presentation(io.BytesIO(file_bytes))

    for slide in presentation.slides:
        for shape in slide.shapes:
            if not hasattr(shape, "text"):
                continue
            for match in PLACEHOLDER_PATTERN.findall(shape.text):
                placeholders.add(match)

    return placeholders


def replace_in_shape_text(shape, replacement_map: dict[str, str]) -> None:
    if not hasattr(shape, "text_frame") or shape.text_frame is None:
        return

    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            original = run.text
            updated = original
            for key, value in replacement_map.items():
                updated = re.sub(r"{{\s*" + re.escape(key) + r"\s*}}", value, updated)
            run.text = updated


def replace_placeholders_in_pptx(file_bytes: bytes, replacement_map: dict[str, str]) -> bytes:
    presentation = Presentation(io.BytesIO(file_bytes))

    for slide in presentation.slides:
        for shape in slide.shapes:
            replace_in_shape_text(shape, replacement_map)

    output_stream = io.BytesIO()
    presentation.save(output_stream)
    output_stream.seek(0)
    return output_stream.getvalue()


def safe_filename(value: str) -> str:
    cleaned = re.sub(r"[^a-zA-Z0-9_-]+", "_", value.strip())
    return cleaned[:80] or "participant"


def build_bulk_output(template_bytes: bytes, participants_df: pd.DataFrame) -> bytes:
    output_stream = io.BytesIO()

    with zipfile.ZipFile(output_stream, mode="w", compression=zipfile.ZIP_DEFLATED) as zip_output:
        manifest_rows: list[dict[str, str]] = []

        for idx, row in participants_df.iterrows():
            replacement_map = {
                column: "" if pd.isna(value) else str(value)
                for column, value in row.to_dict().items()
            }

            generated_ppt = replace_placeholders_in_pptx(template_bytes, replacement_map)

            identity = replacement_map.get("student_name") or replacement_map.get("name") or f"row_{idx + 1}"
            filename = f"certificate_{idx + 1:03d}_{safe_filename(identity)}.pptx"
            zip_output.writestr(filename, generated_ppt)

            manifest_rows.append({"certificate_file": filename, **replacement_map})

        manifest_df = pd.DataFrame(manifest_rows)
        manifest_excel = io.BytesIO()
        manifest_df.to_excel(manifest_excel, index=False)
        manifest_excel.seek(0)
        zip_output.writestr("certificates_manifest.xlsx", manifest_excel.getvalue())

    output_stream.seek(0)
    return output_stream.getvalue()


@app.get("/")
def index():
    return render_template("index.html")


@app.post("/generate")
def generate_certificates():
    template_file = request.files.get("template_file")
    participants_file = request.files.get("participants_file")

    if not template_file or not participants_file:
        flash("Please upload both a PPT template and participant Excel file.")
        return redirect(url_for("index"))

    if not template_file.filename.lower().endswith(".pptx"):
        flash("Template must be a .pptx file.")
        return redirect(url_for("index"))

    if not participants_file.filename.lower().endswith((".xlsx", ".xls")):
        flash("Participants file must be .xlsx or .xls.")
        return redirect(url_for("index"))

    template_bytes = template_file.read()

    try:
        participants_df = pd.read_excel(participants_file)
    except Exception as exc:  # noqa: BLE001
        flash(f"Unable to read participant Excel file: {exc}")
        return redirect(url_for("index"))

    if participants_df.empty:
        flash("Participant Excel has no rows.")
        return redirect(url_for("index"))

    placeholders = extract_placeholders_from_pptx(template_bytes)
    missing_columns = [placeholder for placeholder in placeholders if placeholder not in participants_df.columns]

    if missing_columns:
        flash(
            "Missing required columns in Excel for placeholders: " + ", ".join(sorted(missing_columns))
        )
        return redirect(url_for("index"))

    zip_bytes = build_bulk_output(template_bytes, participants_df)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

    return send_file(
        io.BytesIO(zip_bytes),
        as_attachment=True,
        download_name=f"certificates_{timestamp}.zip",
        mimetype="application/zip",
    )


if __name__ == "__main__":
    app.run(debug=True, host="0.0.0.0", port=5000)
